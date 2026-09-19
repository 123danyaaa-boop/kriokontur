"""D-07, D-13, D-21: сводные представления одного прогона.

KPI по худшему году, сравнение сценариев, разложение эффекта обязательного стресса и полный
набор разделов выгрузки. Проверяется главное свойство: интерфейс, выгрузка и записка берут
числа из одного расчёта, а не пересчитывают их по-своему.
"""
import csv
import io

import pytest

from kriokontur import export
from kriokontur.caseinput import load_case
from kriokontur.engine import run
from kriokontur.plan import Plan
from kriokontur.reporting import (export_sections, kpi_rows, scenario_comparison,
                                  stress_decomposition, strategy_comparison)
from kriokontur.risks import load_risks
from kriokontur.scenarios import load_all


@pytest.fixture(scope="module")
def setup():
    case, scen = load_case(), load_all()
    plan = Plan.load("configs/plans/final-candidate.json")
    return case, scen, plan


def kpi(rows, kpi_id):
    return next(r for r in rows if r["kpi_id"] == kpi_id)


def test_kpi_reports_worst_year_not_average(setup):
    """D-21. Средняя доля обслуживания скрывает провал года, ограничение проверяется ежегодно."""
    case, scen, _ = setup
    thin = Plan.load("configs/plans/final-candidate.json")
    thin.reservations["A"] = {**thin.reservations["A"], 2036: 40.0}   # искусственно рвём 2036 год
    res = run(case, scen["BASE"], thin)
    rows = kpi_rows(case, res)
    worst = kpi(rows, "SL_TOTAL_WORST")
    by_year = {y.year: y.sl_total for y in res.years}
    assert worst["value"] == pytest.approx(min(by_year.values()), abs=1e-12)
    assert worst["period"] == min(by_year, key=by_year.get) == 2036
    assert worst["status"] == "fail"
    assert worst["value"] < res.totals["sl_total"]        # среднее выглядит лучше худшего года


def test_kpi_matches_the_run_it_describes(setup):
    case, scen, plan = setup
    res = run(case, scen["MANDATORY_STRESS"], plan)
    rows = kpi_rows(case, res)
    assert kpi(rows, "PV")["value"] == pytest.approx(res.totals["discounted_cost_mln"], abs=1e-9)
    assert kpi(rows, "TOTAL")["value"] == pytest.approx(res.totals["total_cost_mln"], abs=1e-9)
    assert kpi(rows, "SHORTAGE")["value"] == pytest.approx(res.totals["shortage_t"], abs=1e-9)
    assert kpi(rows, "CAPEX_2040")["value"] == pytest.approx(res.totals["capex"], abs=1e-9)
    assert all(r["formula"] for r in rows), "у каждого KPI должна быть формула"


def test_scenario_comparison_repeats_direct_runs(setup):
    case, scen, plan = setup
    rows = scenario_comparison(case, scen, plan)
    assert [r["scenario_id"] for r in rows] == ["BASE", "MANDATORY_STRESS", "LOW_DEMAND", "HIGH_DEMAND"]
    for row in rows:
        res = run(case, scen[row["scenario_id"]], plan)
        assert row["pv_mln"] == pytest.approx(res.totals["discounted_cost_mln"], abs=1e-9)
        assert row["shortage_t"] == pytest.approx(res.totals["shortage_t"], abs=1e-9)
        assert row["sl_total_worst"] == pytest.approx(min(y.sl_total for y in res.years), abs=1e-12)
    base = rows[0]["pv_mln"]
    assert rows[1]["delta_pv_mln"] == pytest.approx(rows[1]["pv_mln"] - base, abs=1e-9)


def test_strategy_comparison_covers_every_plan(setup):
    case, scen, plan = setup
    plans = {"final-candidate": plan, "no-lunar": Plan.load("configs/plans/no-lunar.json")}
    rows = strategy_comparison(case, scen, plans, ["BASE", "MANDATORY_STRESS"])
    assert {r["plan_id"] for r in rows} == {"final-candidate", "no-lunar"}
    assert len(rows) == 4


def test_stress_decomposition_adds_up(setup):
    """D-13. Сумма одиночных эффектов плюс совместный равна полному эффекту стресса."""
    case, scen, plan = setup
    dec = stress_decomposition(case, scen, plan)
    singles = sum(r["delta_pv_mln"] for r in dec["components"])
    assert singles + dec["interaction_pv_mln"] == pytest.approx(dec["total_delta_pv_mln"], abs=1e-9)
    assert dec["stress_pv_mln"] - dec["base_pv_mln"] == pytest.approx(dec["total_delta_pv_mln"], abs=1e-9)
    by_key = {r["component"]: r for r in dec["components"]}
    assert by_key["demand"]["delta_pv_mln"] > 0
    assert by_key["prices"]["delta_pv_mln"] > 0
    # в финальном плане Луны нет, поэтому её недопоставка ничего не меняет
    assert by_key["isru_delivery"]["delta_pv_mln"] == pytest.approx(0.0, abs=1e-9)


def test_stress_decomposition_sees_isru_when_the_plan_has_it(setup):
    """Разложение обязано реагировать на состав плана, а не выдавать константу."""
    case, scen, _ = setup
    dec = stress_decomposition(case, scen, Plan.load("configs/plans/robust.json"))
    by_key = {r["component"]: r for r in dec["components"]}
    assert by_key["isru_delivery"]["delta_pv_mln"] > 0


def test_export_has_every_required_section(setup):
    """D-07. В выгрузке есть исходные данные, KPI, сравнение сценариев, разложение и риски."""
    case, scen, plan = setup
    res = run(case, scen["BASE"], plan)
    extras = export_sections(case, scen, plan, res, risks=load_risks())
    text = export.to_csv(case, res, scen["BASE"], extras)
    sections, cur = {}, None
    for row in csv.reader(io.StringIO(text), delimiter=";"):
        if len(row) == 1 and row[0]:
            cur = row[0]
            sections[cur] = []
            continue
        if cur and row:
            sections[cur].append(row)
    for name in ("yearly_balance", "source_schedule", "financial_breakdown", "inventory_trace",
                 "constraint_checks", "input_data", "kpi", "scenario_comparison",
                 "stress_decomposition", "risk_register"):
        assert name in sections and len(sections[name]) > 1, name
    # KPI в выгрузке совпадает с KPI прогона
    pv_row = next(r for r in sections["kpi"] if r[0] == "PV")
    assert float(pv_row[3]) == pytest.approx(res.totals["discounted_cost_mln"], abs=0.01)
    # сравнение сценариев содержит оба контрольных сценария
    ids = {r[0] for r in sections["scenario_comparison"][1:]}
    assert {"BASE", "MANDATORY_STRESS"} <= ids


def test_xlsx_has_a_sheet_per_section(setup):
    from openpyxl import load_workbook

    from kriokontur import paths
    case, scen, plan = setup
    res = run(case, scen["BASE"], plan)
    extras = export_sections(case, scen, plan, res, risks=load_risks())
    target = paths.RESULTS / "test-sections.xlsx"
    try:
        path = export.write_xlsx(case, res, scen["BASE"], target, extras)
        wb = load_workbook(path, read_only=True)
        for name in ("yearly_balance", "kpi", "scenario_comparison", "stress_decomposition",
                     "risk_register", "input_data"):
            assert name in wb.sheetnames, wb.sheetnames
        wb.close()
    finally:
        target.unlink(missing_ok=True)
