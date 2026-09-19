"""D-05, D-06, D-22, D-24, D-26: экономика периода и согласованность выгрузки.

D-05  OPEX инвестиции начисляется за месяцы работы объекта, а не за весь год финансирования.
D-06  плата за резерв в выгрузке берётся из движка вместе с долей года.
D-22  XLSX хранит числа числами.
D-24  момент ввода модернизации хранилища задан явным параметром.
D-26  плата за хранение считается по среднему запасу с учётом времени.
"""
import io
import json
from pathlib import Path

import pytest

from kriokontur import export
from kriokontur.caseinput import load_case
from kriokontur.engine import availability, months_available, run
from kriokontur.plan import Plan
from kriokontur.scenarios import load_all

ISRU_PLAN = {"ZBO": 2036, "LUNAR_ISRU": 2036, "EARTH_NEW_OPTION": 2035, "EARTH_NEW_EXERCISE": 2036}


@pytest.fixture(scope="module")
def case_and_scenarios():
    return load_case(), load_all()


def test_isru_opex_is_prorated_by_months_in_operation(case_and_scenarios):
    """D-05. Луна вводится в марте 2038, значит за 2038 год OPEX = 70 × 10/12, а не 70."""
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/robust.json")
    assert plan.investments["LUNAR_ISRU"] is not None
    res = run(case, scen["BASE"], plan)
    avail = availability(case, plan, scen["BASE"])
    zbo_opex = case.storage["ZBO"].fixed_opex_mln_per_year
    isru_opex = case.investments["LUNAR_ISRU"].fixed_opex_mln_per_year
    for i, y in enumerate(res.years):
        fraction = months_available(avail["D"], i) / 12
        expected = (zbo_opex if plan.investments["ZBO"] is not None and y.year >= plan.investments["ZBO"] else 0.0)
        expected += isru_opex * fraction
        assert y.cost["fixed_opex"] == pytest.approx(expected, abs=1e-9), y.year
    first_year = next(y for y in res.years if y.cost["fixed_opex"] > zbo_opex)
    assert first_year.year == 2038
    assert first_year.cost["fixed_opex"] == pytest.approx(zbo_opex + isru_opex * 10 / 12, abs=1e-9)


def test_opex_is_zero_before_commissioning(case_and_scenarios):
    """D-05. Финансирование в 2036 не создаёт OPEX до физического ввода канала."""
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/robust.json")
    res = run(case, scen["BASE"], plan)
    zbo_opex = case.storage["ZBO"].fixed_opex_mln_per_year
    for y in res.years:
        if y.year < 2038:
            assert y.cost["fixed_opex"] <= zbo_opex + 1e-9, y.year


def test_holding_cost_uses_time_weighted_average_stock(case_and_scenarios):
    """D-26. База платы за хранение — середина отрезка «после поступления → конец месяца»."""
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/final-candidate.json")
    res = run(case, scen["BASE"], plan)
    for y in res.years:
        rate = case.storage["ZBO" if y.year >= 2036 else "BASE"].holding_cost_mln_per_t_year
        months = [m for m in res.months if m.year == y.year]
        expected = sum(m.avg_stock_t for m in months) / 12 * rate
        assert y.cost["holding"] == pytest.approx(expected, abs=1e-9), y.year
        assert y.avg_stock_t == pytest.approx(sum(m.avg_stock_t for m in months) / 12, abs=1e-9)
    # выбранная конвенция строже, чем счёт по остатку на конец месяца: она не занижает расходы
    by_closing = sum(sum(m.closing_t for m in res.months if m.year == y.year) / 12
                     * case.storage["ZBO" if y.year >= 2036 else "BASE"].holding_cost_mln_per_t_year
                     for y in res.years)
    assert res.totals["holding"] > by_closing


def test_storage_commissioning_lag_is_an_explicit_assumption(case_and_scenarios):
    """D-24. Момент ввода ZBO задаётся параметром: лаг 6 месяцев сдвигает и потери, и ёмкость."""
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/final-candidate.json")
    base = run(case, scen["BASE"], plan)
    plan.assumptions["storage_commissioning_lag_months"] = 6
    late = run(case, scen["BASE"], plan)
    early_months = [m for m in base.months if m.year == 2036]
    late_months = [m for m in late.months if m.year == 2036]
    assert early_months[0].capacity_t == 120 and late_months[0].capacity_t == 70
    assert late_months[6].capacity_t == 120
    assert late.year(2036).losses_t > base.year(2036).losses_t


def test_source_schedule_matches_financial_breakdown(case_and_scenarios):
    """D-06. Сумма платежей по каналам = итогам прогона для всех планов, включая неполные годы."""
    case, scen = case_and_scenarios
    for name in ("final-candidate", "robust", "robust-plus", "no-lunar", "base-only"):
        plan = Plan.load(f"configs/plans/{name}.json")
        for sid in ("BASE", "MANDATORY_STRESS"):
            res = run(case, scen[sid], plan)
            res_sum = sum(sum(y.reservation_payment_mln.values()) for y in res.years)
            var_sum = sum(sum(y.variable_payment_mln.values()) for y in res.years)
            assert res_sum == pytest.approx(res.totals["reservation"], abs=1e-9), f"{name} × {sid}"
            opening_cost = sum(y.opening_stock_cost_mln for y in res.years)
            assert var_sum + opening_cost == pytest.approx(res.totals["procurement"], abs=1e-9), f"{name} × {sid}"


def test_csv_export_repeats_run_numbers(case_and_scenarios):
    """D-06/EXP-02. Выгрузка не пересчитывает ничего: суммы совпадают с прогоном."""
    import csv
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/robust.json")     # у C есть неполный год с резервом
    res = run(case, scen["BASE"], plan)
    text = export.to_csv(case, res, scen["BASE"])
    sections, cur = {}, None
    for row in csv.reader(io.StringIO(text), delimiter=";"):
        if len(row) == 1 and row[0]:
            cur = row[0]
            sections[cur] = []
            continue
        if cur and row:
            sections[cur].append(row)
    src = sections["source_schedule"][1:]
    assert sum(float(r[8]) for r in src) == pytest.approx(res.totals["reservation"], abs=0.01)
    assert sum(float(r[7]) for r in src) + sum(y.opening_stock_cost_mln for y in res.years) == \
        pytest.approx(res.totals["procurement"], abs=0.01)
    fin_total = sections["financial_breakdown"][-1]
    assert float(fin_total[6]) == pytest.approx(res.totals["total_cost_mln"], abs=0.01)
    assert float(fin_total[7]) == pytest.approx(res.totals["discounted_cost_mln"], abs=0.01)


def test_xlsx_keeps_numbers_numeric(case_and_scenarios):
    """D-22. Жюри должно уметь считать прямо в книге, а не парсить строки."""
    from openpyxl import load_workbook

    from kriokontur import paths
    case, scen = case_and_scenarios
    plan = Plan.load("configs/plans/final-candidate.json")
    res = run(case, scen["MANDATORY_STRESS"], plan)
    # запись только внутрь проекта: paths.resolve_for_write намеренно запрещает выход наружу
    target = paths.RESULTS / "test-xlsx-types.xlsx"
    try:
        path = export.write_xlsx(case, res, scen["MANDATORY_STRESS"], target)
        wb = load_workbook(path, read_only=True)
        rows = list(wb["financial_breakdown"].iter_rows(values_only=True))
        assert all(isinstance(v, (int, float)) for v in rows[1][1:]), rows[1]
        assert rows[-1][7] == pytest.approx(res.totals["discounted_cost_mln"], abs=0.01)
        balance = list(wb["yearly_balance"].iter_rows(values_only=True))
        assert isinstance(balance[1][1], (int, float))
        wb.close()
    finally:
        Path(target).unlink(missing_ok=True)
